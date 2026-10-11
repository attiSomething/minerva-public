"""Turning files from other programs into plain-text drafts.

Word (.docx, which is also how Google Docs downloads), OpenDocument (.odt),
Rich Text (.rtf), Markdown (.md), web pages (.html), plain text and PDFs
with a text layer. Line breaks are kept exactly where the writer put them,
which matters for poems; formatting is dropped. Anything that couldn't be
carried over (images, tables, footnotes) is noted so the writer knows.

The files are only read. Creating the entries is intake.create_entries.
"""

import html
import html.parser
import io
import os
import re
import zipfile
import xml.etree.ElementTree as ET

from . import intake

MAX_FILE_BYTES = 25 * 1024 * 1024
EXTENSIONS = (".docx", ".odt", ".rtf", ".md", ".markdown", ".txt", ".text", ".html", ".htm", ".pdf")


class ImportFileError(ValueError):
    user_facing = True


def _decode(data):
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def _normalize(text):
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\u00a0", " ").replace("\u2028", "\n")
    lines = [ln.rstrip() for ln in text.split("\n")]
    text = "\n".join(lines).strip("\n")
    return re.sub(r"\n{4,}", "\n\n\n", text)


def _join_paragraphs(paras):
    """Paragraphs to text. A document that already spaces its paragraphs
    with empty ones keeps them as they are; otherwise long paragraphs
    (prose) get a blank line between them and short ones (lines of a poem)
    don't."""
    if not paras:
        return ""
    has_blank = any(not p.strip() for p in paras)
    filled = [p for p in paras if p.strip()]
    avg = sum(len(p) for p in filled) / max(1, len(filled))
    sep = "\n" if has_blank or avg < 100 else "\n\n"
    return sep.join(paras)


# ---- Word --------------------------------------------------------------------

def _docx(data):
    import docx

    notes = []
    try:
        d = docx.Document(io.BytesIO(data))
    except Exception:
        raise ImportFileError("This Word file couldn't be opened. It may be damaged, or an older .doc file: "
                              "open it in Word and save it as .docx.") from None
    title = None
    paras = []
    for p in d.paragraphs:
        style = (p.style.name if p.style is not None else "") or ""
        text = p.text
        if title is None and not paras and text.strip() and (style == "Title" or style.startswith("Heading 1")):
            title = text.strip()
            continue
        if title is None and not any(x.strip() for x in paras) and not text.strip():
            continue
        paras.append(text)
    if d.tables:
        notes.append(f"{len(d.tables)} table(s) left out")
    body = d.element.body
    ns = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    pics = len(body.findall(f".//{ns}blip"))
    if pics:
        notes.append(f"{pics} image(s) left out")
    if any(rel.reltype.endswith("/footnotes") for rel in d.part.rels.values()):
        notes.append("footnotes left out")
    if not title and d.core_properties.title:
        title = d.core_properties.title.strip()
    return _join_paragraphs(paras), title, notes


# ---- OpenDocument ------------------------------------------------------------

_ODT = "{urn:oasis:names:tc:opendocument:xmlns:text:1.0}"


def _odt_text(el):
    out = [el.text or ""]
    for child in el:
        tag = child.tag
        if tag == f"{_ODT}line-break":
            out.append("\n")
        elif tag == f"{_ODT}tab":
            out.append("\t")
        elif tag == f"{_ODT}s":
            out.append(" " * int(child.get(f"{_ODT}c", "1")))
        elif tag in (f"{_ODT}note",):
            pass
        else:
            out.append(_odt_text(child))
        out.append(child.tail or "")
    return "".join(out)


def _odt(data):
    notes = []
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            root = ET.fromstring(z.read("content.xml"))
    except Exception:
        raise ImportFileError("This OpenDocument file couldn't be opened.") from None
    title, paras = None, []
    for el in root.iter():
        if el.tag == f"{_ODT}h" and title is None and not any(p.strip() for p in paras):
            title = _odt_text(el).strip() or None
        elif el.tag == f"{_ODT}p":
            paras.append(_odt_text(el))
    if root.find(".//{urn:oasis:names:tc:opendocument:xmlns:table:1.0}table") is not None:
        notes.append("tables left out")
    if root.find(".//{urn:oasis:names:tc:opendocument:xmlns:drawing:1.0}image") is not None:
        notes.append("images left out")
    while paras and not paras[0].strip():
        paras.pop(0)
    return _join_paragraphs(paras), title, notes


# ---- RTF ---------------------------------------------------------------------

_RTF_SKIP = {"fonttbl", "colortbl", "stylesheet", "info", "pict", "header", "footer", "headerl", "headerr",
             "footerl", "footerr", "footnote", "field", "fldinst", "xmlnstbl", "listtable", "listoverridetable",
             "rsidtbl", "generator", "themedata", "colorschememapping", "latentstyles", "datastore", "object"}


def rtf_to_text(rtf):
    """Plain text from RTF: paragraphs and line breaks kept, the rest
    dropped. Handles the escapes Word and WordPad write."""
    out = []
    stack = []
    skip = False
    uc = 1
    pending_skip = 0
    i, n = 0, len(rtf)
    while i < n:
        c = rtf[i]
        if c == "{":
            stack.append((skip, uc))
            i += 1
            continue
        if c == "}":
            skip, uc = stack.pop() if stack else (False, 1)
            i += 1
            continue
        if c == "\\":
            m = re.match(r"\\([a-zA-Z]+)(-?\d+)? ?|\\'([0-9a-fA-F]{2})|\\(.)", rtf[i:], re.S)
            if not m:
                i += 1
                continue
            i += m.end()
            word, arg, hexv, sym = m.groups()
            if hexv:
                if pending_skip:
                    pending_skip -= 1
                elif not skip:
                    out.append(bytes([int(hexv, 16)]).decode("cp1252", errors="replace"))
                continue
            if sym is not None:
                if sym == "*":
                    skip = True
                elif sym in "\\{}" and not skip:
                    out.append(sym)
                elif sym == "~" and not skip:
                    out.append(" ")
                elif sym in "\n\r" and not skip:
                    out.append("\n")
                continue
            if word in _RTF_SKIP:
                skip = True
            elif skip:
                continue
            elif word in ("par", "line"):
                out.append("\n")
            elif word == "tab":
                out.append("\t")
            elif word == "uc":
                uc = int(arg or 1)
            elif word == "u":
                v = int(arg)
                out.append(chr(v + 65536 if v < 0 else v))
                pending_skip = uc
            elif word in ("emdash",):
                out.append("\u2014")
            elif word in ("endash",):
                out.append("\u2013")
            elif word in ("lquote", "rquote"):
                out.append("\u2018" if word == "lquote" else "\u2019")
            elif word in ("ldblquote", "rdblquote"):
                out.append("\u201c" if word == "ldblquote" else "\u201d")
            continue
        if c in "\r\n":
            i += 1
            continue
        if pending_skip:
            pending_skip -= 1
        elif not skip:
            out.append(c)
        i += 1
    return "".join(out)


def _rtf(data):
    text = _decode(data)
    if not text.lstrip().startswith("{\\rtf"):
        raise ImportFileError("This doesn't look like a Rich Text file.")
    notes = ["images left out"] if "\\pict" in text else []
    return rtf_to_text(text), None, notes


# ---- Markdown ----------------------------------------------------------------

FRONT_MATTER_RE = re.compile(r"\A---\s*\n(.*?)\n(?:---|\.\.\.)\s*\n", re.S)


def front_matter(text):
    """(fields, rest) from Markdown with a YAML front-matter block. Only
    simple 'key: value' and 'key: [a, b]' / '- item' lists are read."""
    m = FRONT_MATTER_RE.match(text)
    if not m:
        return {}, text
    fields, key = {}, None
    for line in m.group(1).split("\n"):
        item = re.match(r"^\s*-\s+(.*)$", line)
        if item and key:
            fields.setdefault(key, [])
            if isinstance(fields[key], list):
                fields[key].append(item.group(1).strip().strip("'\""))
            continue
        kv = re.match(r"^([A-Za-z0-9_-]+)\s*:\s*(.*)$", line)
        if kv:
            key, val = kv.group(1).lower(), kv.group(2).strip()
            if val.startswith("[") and val.endswith("]"):
                fields[key] = [v.strip().strip("'\"") for v in val[1:-1].split(",") if v.strip()]
            elif val:
                fields[key] = val.strip("'\"")
            else:
                fields[key] = []
    return fields, text[m.end():]


def markdown_to_text(text):
    out = []
    in_code = False
    for line in text.split("\n"):
        if line.strip().startswith("```"):
            in_code = not in_code
            continue
        if in_code:
            out.append(line)
            continue
        had = line.strip()
        line = re.sub(r"!\[\[[^\]]*\]\]|!\[[^\]]*\]\([^)]*\)", "", line)        # images
        if had and not line.strip():
            continue
        line = re.sub(r"\[\[([^\]|]+)\|([^\]]+)\]\]", r"\2", line)              # [[link|alias]]
        line = re.sub(r"\[\[([^\]]+)\]\]", r"\1", line)                          # [[link]]
        line = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", line)                     # [text](url)
        line = re.sub(r"^\s{0,3}#{1,6}\s+", "", line)                             # headings
        line = re.sub(r"^\s{0,3}>\s?", "", line)                                  # quotes
        line = re.sub(r"(\*\*|__)(?=\S)(.+?)(?<=\S)\1", r"\2", line)              # bold
        line = re.sub(r"(?<![\w*])\*(?=\S)(.+?)(?<=\S)\*(?![\w*])", r"\1", line)  # *italic*
        line = re.sub(r"(?<![\w_])_(?=\S)(.+?)(?<=\S)_(?![\w_])", r"\1", line)    # _italic_
        line = re.sub(r"~~(.+?)~~", r"\1", line)
        line = re.sub(r"`([^`]+)`", r"\1", line)
        line = re.sub(r"(\\| {2,})$", "", line)                                   # hard-break markers
        line = re.sub(r"\\([\\`*_{}\[\]()#+\-.!>])", r"\1", line)                 # escapes
        out.append(line)
    return "\n".join(out)


def _markdown(data):
    raw = _decode(data).replace("\r\n", "\n")
    fields, rest = front_matter(raw)
    title = fields.get("title") if isinstance(fields.get("title"), str) else None
    lines = rest.split("\n")
    while lines and not lines[0].strip():
        lines.pop(0)
    if lines and re.match(r"^#\s+\S", lines[0]):
        title = title or re.sub(r"^#\s+", "", lines[0]).strip()
        lines.pop(0)
    notes = ["images left out"] if re.search(r"!\[", rest) else []
    return markdown_to_text("\n".join(lines)), title, notes, fields


# ---- HTML --------------------------------------------------------------------

class _HTMLText(html.parser.HTMLParser):
    BLOCK = {"p", "div", "h1", "h2", "h3", "h4", "h5", "h6", "li", "tr", "blockquote", "pre", "section", "article"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.out, self.skip, self.title, self.in_title, self.images = [], 0, None, False, 0

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style", "head"):
            self.skip += 1
        if tag == "title":
            self.in_title = True
        if tag == "br":
            self.out.append("\n")
        if tag == "img":
            self.images += 1
        if tag in self.BLOCK:
            self.out.append("\n\x00")

    def handle_endtag(self, tag):
        if tag in ("script", "style", "head"):
            self.skip = max(0, self.skip - 1)
        if tag == "title":
            self.in_title = False

    def handle_data(self, data):
        if self.in_title:
            self.title = (self.title or "") + data
        elif not self.skip:
            self.out.append(re.sub(r"\s+", " ", data))


def _html(data):
    p = _HTMLText()
    p.feed(_decode(data))
    text = "".join(p.out)
    paras = [x.strip(" ") for x in text.split("\n\x00")]
    paras = [x for x in paras]
    while paras and not paras[0].strip():
        paras.pop(0)
    notes = [f"{p.images} image(s) left out"] if p.images else []
    title = html.unescape(p.title).strip() if p.title else None
    return _join_paragraphs([x.replace("\x00", "") for x in paras]), title, notes


# ---- PDF ---------------------------------------------------------------------

def _pdf(data):
    import pypdfium2 as pdfium

    try:
        pdf = pdfium.PdfDocument(data)
    except Exception:
        raise ImportFileError("This PDF couldn't be opened. It may be damaged or password-protected.") from None
    try:
        pages = []
        for i in range(len(pdf)):
            tp = pdf[i].get_textpage()
            pages.append(tp.get_text_range())
        title = (pdf.get_metadata_dict().get("Title") or "").strip() or None
    finally:
        pdf.close()
    text = "\n\n".join(p.strip("\r\n") for p in pages)
    if len(re.sub(r"\s", "", text)) < 20:
        raise ImportFileError("This PDF has no text in it, only page images. Bring it in through Scans "
                              "so the pages can be read.")
    return text, title, ["page layout not kept: check line breaks"]


# ---- entry point -------------------------------------------------------------

def _useful_title(title):
    """A title from a file's own properties, unless it's a placeholder a
    program filled in."""
    t = re.sub(r"^Microsoft Word - ", "", (title or "").strip())
    if not t or t.lower() in ("untitled", "document", "untitled document") or re.search(r"\.\w{2,5}$", t):
        return None
    return t


def _drop_title_line(text, title):
    """A draft whose first line is just its title, then a blank line: the
    title is already the entry's title, so it isn't repeated in the draft
    (it would print twice in a manuscript)."""
    key = lambda t: re.sub(r"[\W_]+", " ", t).strip().casefold()  # noqa: E731
    lines = text.split("\n")
    i = next((k for k, line in enumerate(lines) if line.strip()), None)
    if i is None or key(lines[i]) != key(title) or not key(title):
        return text
    rest = lines[i + 1:]
    if not rest or rest[0].strip() or not any(x.strip() for x in rest):
        return text
    return "\n".join(rest).lstrip("\n")


def convert(name, data):
    """{"name", "title", "text", "words", "notes", "fields"} for one file,
    or raise ImportFileError."""
    ext = os.path.splitext(name)[1].lower()
    if ext not in EXTENSIONS:
        raise ImportFileError("Minerva can't read this kind of file. It takes Word, OpenDocument, Rich Text, "
                              "Markdown, web pages, text and PDF files.")
    if len(data) > MAX_FILE_BYTES:
        raise ImportFileError("This file is too large to bring in as one entry.")
    fields = {}
    if ext == ".docx":
        text, title, notes = _docx(data)
    elif ext == ".odt":
        text, title, notes = _odt(data)
    elif ext == ".rtf":
        text, title, notes = _rtf(data)
    elif ext in (".md", ".markdown"):
        text, title, notes, fields = _markdown(data)
    elif ext in (".html", ".htm"):
        text, title, notes = _html(data)
    elif ext == ".pdf":
        text, title, notes = _pdf(data)
    else:
        text, title, notes = _decode(data), None, []
    text = _normalize(text)
    stem = os.path.splitext(os.path.basename(name.replace("\\", "/")))[0]
    title = intake.clean_title(_useful_title(title) or stem) or intake.clean_title(stem) or "Untitled"
    text = _drop_title_line(text, title)
    return {"name": name, "title": title, "text": text + "\n" if text else "",
            "words": len(text.split()), "notes": notes, "fields": fields}
