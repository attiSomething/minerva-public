"""Standard manuscript layout: the plain format most magazines and contests
ask poets to send.

Each piece starts on a new US Letter page with the writer's contact details
at the top left and its line count at the top right, then the title in
capitals, then the text in 12 pt Times New Roman, single-spaced, with a
blank line between stanzas. A stanza is kept on one page when it fits.
Later pages carry a running header: surname / title / page.

The contact details are kept in Minerva's own settings on this computer
(never in the library) and filled in under Settings > Manuscript details.
"""

import re

from . import config as app_config

CONTACT_FIELDS = ("name", "address", "email", "phone", "website")
INCH = 72.0
FONT_SIZE = 12


def contact():
    saved = app_config.load().get("contact") or {}
    return {k: str(saved.get(k) or "") for k in CONTACT_FIELDS}


def save_contact(values):
    cfg = app_config.load()
    out = {}
    for k in CONTACT_FIELDS:
        v = str((values or {}).get(k) or "").replace("\r\n", "\n").strip()
        out[k] = v[:400] if k == "address" else " ".join(v.split())[:200]
    cfg["contact"] = out
    app_config.save(cfg)
    return out


def contact_lines(c, fallback_name=""):
    lines = [c["name"] or fallback_name]
    lines += [ln.strip() for ln in c["address"].split("\n") if ln.strip()]
    lines += [c[k] for k in ("phone", "email", "website") if c[k]]
    return [ln for ln in lines if ln]


def surname(name):
    parts = (name or "").split()
    return parts[-1] if parts else ""


def blocks(body):
    """Stanzas, each a list of lines."""
    return [b.split("\n") for b in re.split(r"\n\s*\n", body.strip("\n")) if b.strip()]


def line_count(body):
    return sum(1 for ln in body.split("\n") if ln.strip())


def count_label(body):
    n = line_count(body)
    return f"{n} line{'' if n == 1 else 's'}"


def pieces_label(doc):
    n = len(doc["pieces"])
    poems = all(p.get("category") == "Poetry" for p in doc["pieces"])
    word = ("poem" if poems else "piece") + ("" if n == 1 else "s")
    return f"{n} {word}"


def running_head(name, title):
    s = surname(name)
    return f"{s} / {title.upper()} / " if s else f"{title.upper()} / "


# ---- plain text ---------------------------------------------------------------

def as_text(doc, c):
    out = []
    contact_block = contact_lines(c, doc.get("author", ""))
    if not doc["single"] and doc.get("title_page"):
        out += contact_block + ["", "", doc["title"].upper()]
        if doc.get("subtitle"):
            out.append(doc["subtitle"])
        out += ["", pieces_label(doc), "", "\f"]
    for i, p in enumerate(doc["pieces"]):
        if i:
            out.append("\f")
        first = contact_block[0] if contact_block else ""
        out.append(f"{first}{' ' * max(2, 60 - len(first) - len(count_label(p['body'])))}{count_label(p['body'])}")
        out += contact_block[1:]
        out += ["", "", p["title"].upper(), "", p["body"].strip("\n"), ""]
    return "\n".join(out).rstrip("\n") + "\n"


def as_text_file(doc, path, c):
    with open(path, "w", encoding="utf-8") as fh:
        fh.write(as_text(doc, c))


# ---- Word ---------------------------------------------------------------------

def _page_field(paragraph):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def _letter(section):
    from docx.shared import Inches
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Inches(1))
    section.header_distance = Inches(0.5)


def as_docx(doc, path, c):
    import docx
    from docx.enum.section import WD_SECTION
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.shared import Inches, Pt

    document = docx.Document()
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(FONT_SIZE)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.line_spacing = 1.0
    contact_block = contact_lines(c, doc.get("author", ""))
    name = contact_block[0] if contact_block else ""

    def blank(n=1):
        for _ in range(n):
            document.add_paragraph()

    def contact_with_count(label):
        p = document.add_paragraph()
        p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
        p.add_run(f"{name}\t{label}" if label else name)
        for line in contact_block[1:]:
            document.add_paragraph(line)

    started = False

    def new_section(head_title):
        nonlocal started
        section = document.add_section(WD_SECTION.NEW_PAGE) if started else document.sections[0]
        started = True
        _letter(section)
        section.different_first_page_header_footer = True
        section.header.is_linked_to_previous = False
        section.first_page_header.is_linked_to_previous = False
        hp = section.header.paragraphs[0]
        hp.text = ""
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hp.add_run(running_head(name, head_title))
        _page_field(hp)

    if not doc["single"] and doc.get("title_page"):
        new_section(doc["title"])
        for line in contact_block:
            document.add_paragraph(line)
        blank(10)
        t = document.add_paragraph(doc["title"].upper())
        t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if doc.get("subtitle"):
            s = document.add_paragraph(doc["subtitle"])
            s.alignment = WD_ALIGN_PARAGRAPH.CENTER
        blank()
        n = document.add_paragraph(pieces_label(doc))
        n.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for p in doc["pieces"]:
        new_section(p["title"])
        contact_with_count(count_label(p["body"]))
        blank(3)
        t = document.add_paragraph(p["title"].upper())
        t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        t.paragraph_format.space_after = Pt(FONT_SIZE * 2)
        t.paragraph_format.keep_with_next = True
        for stanza in blocks(p["body"]):
            para = document.add_paragraph()
            para.paragraph_format.keep_together = True
            para.paragraph_format.space_after = Pt(FONT_SIZE)
            run = para.add_run()
            for i, line in enumerate(stanza):
                if i:
                    run.add_break()
                run.add_text(line)
    document.save(path)


# ---- PDF ----------------------------------------------------------------------

def _times():
    """Times New Roman when Windows has it, else the PDF built-in Times."""
    import os
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    path = os.path.join(os.environ.get("WINDIR", r"C:\Windows"), "Fonts", "times.ttf")
    if os.path.isfile(path):
        if "MinervaTimes" not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont("MinervaTimes", path))
        return "MinervaTimes"
    return "Times-Roman"


def as_pdf(doc, path, c):
    from xml.sax.saxutils import escape

    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import (BaseDocTemplate, Flowable, Frame, KeepTogether, PageBreak,
                                    PageTemplate, Paragraph, Spacer, Table, TableStyle)

    font = _times()
    w, h = 8.5 * INCH, 11 * INCH
    margin = INCH
    lead = FONT_SIZE * 1.2
    body = ParagraphStyle("ms-body", fontName=font, fontSize=FONT_SIZE, leading=lead)
    centered = ParagraphStyle("ms-center", parent=body, alignment=TA_CENTER)
    right = ParagraphStyle("ms-right", parent=body, alignment=2)
    contact_block = contact_lines(c, doc.get("author", ""))
    name = contact_block[0] if contact_block else ""
    state = {"title": doc["title"], "start": 1}

    class PieceStart(Flowable):
        """Marks where a piece begins, so its first page has no header."""
        def __init__(self, title):
            super().__init__()
            self.title = title

        def wrap(self, *_):
            return 0, 0

        def draw(self):
            state["title"] = self.title
            state["start"] = self.canv.getPageNumber()

    def header(canvas, _doc):
        page = canvas.getPageNumber()
        if page == state["start"]:
            return
        canvas.setFont(font, FONT_SIZE)
        canvas.drawRightString(w - margin, h - margin / 2 - FONT_SIZE / 2,
                               f"{running_head(name, state['title'])}{page}")

    def line(text, style=body):
        return Paragraph(escape(text) if text else "&nbsp;", style)

    def markup(ln):
        stripped = ln.lstrip(" ")
        return "&nbsp;" * (len(ln) - len(stripped)) + escape(stripped)

    def contact_table(label):
        rows = [[line(name), line(label, right) if label else ""]]
        rows += [[line(x), ""] for x in contact_block[1:]]
        t = Table(rows, colWidths=[(w - 2 * margin) * 0.7, (w - 2 * margin) * 0.3])
        t.setStyle(TableStyle([("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                               ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
                               ("VALIGN", (0, 0), (-1, -1), "TOP")]))
        return t

    story = []
    if not doc["single"] and doc.get("title_page"):
        story += [PieceStart(doc["title"]), contact_table(""), Spacer(1, lead * 10),
                  line(doc["title"].upper(), centered)]
        if doc.get("subtitle"):
            story.append(line(doc["subtitle"], centered))
        story += [Spacer(1, lead), line(pieces_label(doc), centered), PageBreak()]
    for i, p in enumerate(doc["pieces"]):
        if i:
            story.append(PageBreak())
        story += [PieceStart(p["title"]), contact_table(count_label(p["body"])), Spacer(1, lead * 3)]
        stanzas = blocks(p["body"])
        title = [line(p["title"].upper(), centered), Spacer(1, lead * 2)]
        for j, stanza in enumerate(stanzas):
            para = Paragraph("<br/>".join(markup(ln) for ln in stanza), body)
            parts = ([*title] if j == 0 else []) + [para, Spacer(1, lead)]
            story.append(KeepTogether(parts))
        if not stanzas:
            story += title

    pdf = BaseDocTemplate(path, pagesize=(w, h), title=doc["title"], author=name, creator="Minerva",
                          leftMargin=margin, rightMargin=margin, topMargin=margin, bottomMargin=margin)
    pdf.addPageTemplates([PageTemplate("ms", [Frame(margin, margin, w - 2 * margin, h - 2 * margin, id="ms",
                                                              leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)],
                                       onPageEnd=header)])
    pdf.build(story)
