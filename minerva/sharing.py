"""Sharing: final drafts, and exporting entries and projects as .txt, .docx
or .pdf (and projects as .epub, see epub.py).

A final draft is a copy of a draft kept in the entry's own "Final" folder,
so it is visible in the library and travels with the entry. Marking a new
final moves the old one into the entry's Archive; nothing is deleted.

Exports are written to the library's "Exports" folder under a name that is
never reused, so an export never replaces an earlier one. Templates and
covers the user adds are copied into the hidden .minerva folder.
"""

import datetime
import io
import json
import os
import re
import shutil
import uuid
from xml.sax.saxutils import escape

from . import core

FORMATS = ("txt", "docx", "pdf")
PROJECT_FORMATS = FORMATS + ("epub",)
INCH = 72.0

# name, width and height in inches
TRIM_SIZES = {
    "letter": ("US Letter (8.5 x 11 in)", 8.5, 11.0),
    "a4": ("A4 (210 x 297 mm)", 8.27, 11.69),
    "a5": ("A5 (148 x 210 mm)", 5.83, 8.27),
    "6x9": ("Trade paperback (6 x 9 in)", 6.0, 9.0),
    "5.5x8.5": ("Digest (5.5 x 8.5 in)", 5.5, 8.5),
    "5x8": ("Small paperback (5 x 8 in)", 5.0, 8.0),
}
# a layout rather than a size: US Letter in standard manuscript format
MANUSCRIPT = "manuscript"
MANUSCRIPT_LABEL = "Standard manuscript (US Letter)"
IMAGE_EXTS = {".jpg", ".jpeg", ".png"}
PROJECT_FIELDS = ("subtitle", "author", "copyright", "isbn", "extra", "trim",
                  "title_page", "contents", "dedication", "epigraph", "acknowledgments", "about")
# optional pages: before the contents, and after the last piece
FRONT_PAGES = (("dedication", "Dedication", "dedication"), ("epigraph", "Epigraph", "epigraph"))
BACK_PAGES = (("acknowledgments", "Acknowledgments"), ("about", "About the Author"))


class ShareError(Exception):
    user_facing = True


# --------------------------------------------------------------------------
# Final drafts
# --------------------------------------------------------------------------

def _final_dir(row):
    return os.path.join(core.entry_dir(row), "Final")


def _final_files(row):
    d = _final_dir(row)
    if not os.path.isdir(d):
        return []
    return sorted(f for f in os.listdir(d)
                  if f.lower().endswith(".txt") and os.path.isfile(os.path.join(d, f)))


def _read(path):
    with open(path, "r", encoding="utf-8", errors="ignore") as fh:
        return fh.read()


def get_final(row):
    """The entry's final draft, or None."""
    files = _final_files(row)
    if not files:
        return None
    path = os.path.join(_final_dir(row), files[0])
    text = _read(path)
    return {"name": files[0], "text": text,
            "marked": datetime.datetime.fromtimestamp(os.path.getmtime(path)).isoformat(timespec="seconds"),
            "same_as_current": text == core.read_entry_content(row)}


def _retire_final(row):
    """Move any current final draft into the entry's Archive."""
    adir = os.path.join(core.entry_dir(row), "Archive")
    for f in _final_files(row):
        os.makedirs(adir, exist_ok=True)
        name = core.unique_archive_name(
            adir, f"{row['title']} (final until {datetime.date.today().isoformat()}).txt")
        shutil.move(os.path.join(_final_dir(row), f), os.path.join(adir, name))


def mark_final(conn, entry_id, archive_name=None):
    """Make the current draft (or one of the entry's archived drafts) its
    final draft."""
    row = core.get_entry_row(conn, entry_id)
    if row is None:
        raise ShareError("That entry no longer exists.")
    text = core.read_archive_file(row, archive_name) if archive_name else core.read_entry_content(row)
    _retire_final(row)
    os.makedirs(_final_dir(row), exist_ok=True)
    with open(os.path.join(_final_dir(row), f"{row['title']} (final).txt"), "w", encoding="utf-8") as fh:
        fh.write(text)
    return get_final(row)


def unmark_final(conn, entry_id):
    row = core.get_entry_row(conn, entry_id)
    if row is None:
        raise ShareError("That entry no longer exists.")
    _retire_final(row)


def has_final(row):
    return bool(_final_files(row))


# --------------------------------------------------------------------------
# What gets exported
# --------------------------------------------------------------------------

HEADER_LINE_RE = re.compile(r"^\s*(Source|Source entries|Category|Date|Location|Time)\s*:", re.IGNORECASE)


def strip_header(text):
    """Drop a transcription header ("Source: ...", the title, "Category:",
    "Date:" lines closed by a ==== line) from the top of a draft."""
    lines = text.replace("\r\n", "\n").split("\n")
    for i, line in enumerate(lines[:14]):
        if line.strip() and set(line.strip()) == {"="}:
            if any(HEADER_LINE_RE.match(ln) for ln in lines[:i]):
                return "\n".join(lines[i + 1:]).strip("\n")
            break
    return text.strip("\n")


def entry_piece(row, which="final"):
    """{title, body} for one entry: its final draft when it has one (unless
    the current draft is asked for), else its current draft."""
    final = get_final(row) if which != "current" else None
    text = final["text"] if final else core.read_entry_content(row)
    return {"title": row["title"], "body": strip_header(text), "category": row["top_folder"],
            "from_final": bool(final)}


def _blocks(body):
    """Paragraphs (or stanzas), each a list of lines."""
    return [b.split("\n") for b in re.split(r"\n\s*\n", body.strip("\n")) if b.strip()]


# --------------------------------------------------------------------------
# Files
# --------------------------------------------------------------------------

def exports_dir():
    d = os.path.join(core.ARCHIVE_ROOT, "Exports")
    os.makedirs(d, exist_ok=True)
    return d


def _safe_name(name):
    return core.INVALID_FILENAME_CHARS_RE.sub("", name).strip(" .")[:90] or "Export"


def _unique_path(stem, ext):
    d = exports_dir()
    stem = _safe_name(stem)
    path = os.path.join(d, f"{stem}.{ext}")
    n = 2
    while os.path.exists(path):
        path = os.path.join(d, f"{stem} ({n}).{ext}")
        n += 1
    return path


def export_path_ok(path):
    """True if path is a file inside the Exports folder."""
    d = os.path.normcase(os.path.normpath(exports_dir()))
    p = os.path.normcase(os.path.normpath(path))
    return p.startswith(d + os.sep) and os.path.isfile(p)


# --------------------------------------------------------------------------
# Templates
# --------------------------------------------------------------------------

def _templates_dir():
    d = os.path.join(core.META_DIR, "templates")
    os.makedirs(d, exist_ok=True)
    return d


def _templates_index():
    try:
        with open(os.path.join(_templates_dir(), "templates.json"), "r", encoding="utf-8") as fh:
            return json.load(fh)
    except (OSError, ValueError):
        return []


def _save_templates_index(items):
    with open(os.path.join(_templates_dir(), "templates.json"), "w", encoding="utf-8") as fh:
        json.dump(items, fh, indent=1)


def list_templates():
    return [t for t in _templates_index() if os.path.isfile(os.path.join(_templates_dir(), t["file"]))]


def add_template(filename, data, name=None):
    ext = os.path.splitext(filename or "")[1].lower()
    if ext not in (".docx", ".pdf"):
        raise ShareError("A template has to be a .docx or .pdf file.")
    if ext == ".docx":
        import docx
        try:
            docx.Document(io.BytesIO(data))
        except Exception:
            raise ShareError(f'"{filename}" couldn\'t be opened as a Word document.')
    else:
        _pdf_background(data)  # raises if it isn't a readable PDF
    tid = uuid.uuid4().hex[:10]
    with open(os.path.join(_templates_dir(), tid + ext), "wb") as fh:
        fh.write(data)
    item = {"id": tid, "name": (name or os.path.splitext(filename)[0]).strip()[:80] or "Template",
            "kind": ext[1:], "file": tid + ext, "added": datetime.date.today().isoformat()}
    _save_templates_index(_templates_index() + [item])
    return item


def remove_template(tid):
    items = _templates_index()
    keep = [t for t in items if t["id"] != tid]
    if len(keep) == len(items):
        raise ShareError("That template is already gone.")
    gone = next(t for t in items if t["id"] == tid)
    try:
        os.remove(os.path.join(_templates_dir(), gone["file"]))
    except OSError:
        pass
    _save_templates_index(keep)


def get_template(tid):
    t = next((t for t in list_templates() if t["id"] == tid), None)
    if t is None:
        raise ShareError("That template is no longer there. Choose another one.")
    return t, os.path.join(_templates_dir(), t["file"])


def _pdf_background(data):
    """(width_pt, height_pt, PNG bytes) of a PDF's first page."""
    import pypdfium2 as pdfium
    try:
        pdf = pdfium.PdfDocument(data)
    except Exception:
        raise ShareError("That file couldn't be opened as a PDF.")
    try:
        if len(pdf) < 1:
            raise ShareError("That PDF has no pages.")
        page = pdf[0]
        w, h = page.get_size()
        img = page.render(scale=2.0).to_pil()
    finally:
        pdf.close()
    buf = io.BytesIO()
    img.convert("RGB").save(buf, "PNG")
    return w, h, buf.getvalue()


# --------------------------------------------------------------------------
# Projects: settings and covers
# --------------------------------------------------------------------------

def _ensure_columns(conn):
    cols = [r["name"] for r in conn.execute("PRAGMA table_info(projects)").fetchall()]
    if "share_settings" not in cols:
        conn.execute("ALTER TABLE projects ADD COLUMN share_settings TEXT")
        conn.commit()


def project_settings(conn, project_id):
    _ensure_columns(conn)
    row = conn.execute("SELECT share_settings FROM projects WHERE id = ?", (project_id,)).fetchone()
    if row is None:
        raise ShareError("That project no longer exists.")
    try:
        saved = json.loads(row["share_settings"] or "{}")
    except ValueError:
        saved = {}
    out = {"subtitle": "", "author": "", "copyright": "", "isbn": "", "extra": "",
           "trim": "6x9", "title_page": True, "contents": True,
           "dedication": "", "epigraph": "", "acknowledgments": "", "about": ""}
    if "author" not in saved:
        # until it's set, the author is the name in the manuscript details
        from . import manuscript
        out["author"] = (manuscript.contact().get("name") or "").strip()
    out.update({k: v for k, v in saved.items() if k in PROJECT_FIELDS})
    out["covers"] = {side: bool(_cover_path(project_id, side)) for side in ("front", "back")}
    return out


def save_project_settings(conn, project_id, values):
    current = project_settings(conn, project_id)
    current.pop("covers", None)
    for k in PROJECT_FIELDS:
        if k in values:
            v = values[k]
            current[k] = bool(v) if k in ("title_page", "contents") else str(v or "").replace("\r\n", "\n").strip()
    if current["trim"] not in TRIM_SIZES and current["trim"] != MANUSCRIPT:
        raise ShareError("Choose one of the listed page sizes.")
    conn.execute("UPDATE projects SET share_settings = ? WHERE id = ?", (json.dumps(current), project_id))
    conn.commit()
    return project_settings(conn, project_id)


def _covers_dir(project_id):
    return os.path.join(core.META_DIR, "projects", str(int(project_id)))


def _cover_path(project_id, side):
    d = _covers_dir(project_id)
    if not os.path.isdir(d):
        return None
    for f in os.listdir(d):
        if os.path.splitext(f)[0] == side:
            return os.path.join(d, f)
    return None


def set_cover(project_id, side, filename, data):
    if side not in ("front", "back"):
        raise ShareError("A cover is either the front or the back.")
    ext = os.path.splitext(filename or "")[1].lower()
    if ext not in IMAGE_EXTS:
        raise ShareError("A cover has to be a .jpg or .png image.")
    from PIL import Image
    try:
        Image.open(io.BytesIO(data)).verify()
    except Exception:
        raise ShareError(f'"{filename}" couldn\'t be opened as an image.')
    remove_cover(project_id, side)
    os.makedirs(_covers_dir(project_id), exist_ok=True)
    with open(os.path.join(_covers_dir(project_id), side + ext), "wb") as fh:
        fh.write(data)


def remove_cover(project_id, side):
    p = _cover_path(project_id, side)
    if p:
        os.remove(p)


def cover_file(project_id, side):
    return _cover_path(project_id, side)


# --------------------------------------------------------------------------
# Export
# --------------------------------------------------------------------------

def export_entry(conn, entry_id, fmt, which="final", template_id=None):
    row = core.get_entry_row(conn, entry_id)
    if row is None:
        raise ShareError("That entry no longer exists.")
    if fmt not in FORMATS:
        raise ShareError("Choose .txt, .docx or .pdf.")
    piece = entry_piece(row, which)
    doc = {"title": row["title"], "pieces": [piece], "trim": "letter", "single": True}
    template = None
    if template_id == MANUSCRIPT:
        doc["manuscript"] = True
        template_id = None
    if template_id:
        template, tpath = get_template(template_id)
        if template["kind"] != fmt:
            raise ShareError(f'"{template["name"]}" is a .{template["kind"]} template; '
                             f"choose .{template['kind']} to use it.")
        template = dict(template, path=tpath)
    path = _unique_path(row["title"], fmt)
    _write(doc, fmt, path, template)
    return {"path": path, "name": os.path.basename(path), "from_final": piece["from_final"]}


def export_project(conn, project_id, fmt):
    if fmt not in PROJECT_FORMATS:
        raise ShareError("Choose .txt, .docx, .pdf or .epub.")
    project = core.get_project(conn, project_id)
    settings = project_settings(conn, project_id)
    rows = [core.get_entry_row(conn, e["id"]) for e in project["entries"]]
    pieces = [entry_piece(r) for r in rows if r is not None]
    if not pieces:
        raise ShareError("This project has no entries to export yet.")
    doc = {
        "title": project["name"], "pieces": pieces, "trim": settings["trim"], "single": False,
        "subtitle": settings["subtitle"], "author": settings["author"],
        "copyright": settings["copyright"], "isbn": settings["isbn"], "extra": settings["extra"],
        "title_page": settings["title_page"], "contents": settings["contents"],
        "front": cover_file(project_id, "front"), "back": cover_file(project_id, "back"),
        "front_pages": [(label, settings[k], css) for k, label, css in FRONT_PAGES if settings[k]],
        "back_pages": [(label, settings[k], "") for k, label in BACK_PAGES if settings[k]],
    }
    if fmt == "epub":
        from . import epub
        path = _unique_path(project["name"], "epub")
        tmp = path + ".minerva-tmp"
        try:
            epub.write(doc | {"front_matter_lines": _front_matter_lines(doc)}, tmp)
            os.replace(tmp, path)
        finally:
            if os.path.exists(tmp):
                os.remove(tmp)
        return {"path": path, "name": os.path.basename(path),
                "finals": sum(1 for p in pieces if p["from_final"]), "entries": len(pieces)}
    if settings["trim"] == MANUSCRIPT:
        doc.update(manuscript=True, trim="letter")
        path = _unique_path(f"{project['name']} - manuscript", fmt)
    else:
        label = "proof" if fmt == "pdf" else "manuscript"
        path = _unique_path(f"{project['name']} - {label} ({settings['trim']})", fmt)
    _write(doc, fmt, path, None)
    return {"path": path, "name": os.path.basename(path),
            "finals": sum(1 for p in pieces if p["from_final"]), "entries": len(pieces)}


def _write(doc, fmt, path, template):
    tmp = path + ".minerva-tmp"
    try:
        if doc.get("manuscript"):
            from . import manuscript
            getattr(manuscript, {"txt": "as_text_file", "docx": "as_docx", "pdf": "as_pdf"}[fmt])(
                doc, tmp, manuscript.contact())
        elif fmt == "txt":
            with open(tmp, "w", encoding="utf-8") as fh:
                fh.write(_as_text(doc))
        elif fmt == "docx":
            _as_docx(doc, tmp, template)
        else:
            _as_pdf(doc, tmp, template)
        os.replace(tmp, path)
    finally:
        if os.path.exists(tmp):
            os.remove(tmp)


def _front_matter_lines(doc):
    """The copyright page, as lines."""
    lines = []
    if doc.get("copyright"):
        lines.append(doc["copyright"])
    if doc.get("isbn"):
        lines.append(f"ISBN {doc['isbn']}" if not doc["isbn"].upper().startswith("ISBN") else doc["isbn"])
    if doc.get("extra"):
        lines += [""] + doc["extra"].split("\n")
    return lines


# ---- plain text ----------------------------------------------------------

def _as_text(doc):
    if doc["single"]:
        p = doc["pieces"][0]
        return f"{p['title']}\n\n{p['body']}\n"
    out = []
    if doc["title_page"]:
        out += [doc["title"].upper()]
        if doc["subtitle"]:
            out.append(doc["subtitle"])
        if doc["author"]:
            out += ["", doc["author"]]
        out.append("")
    fm = _front_matter_lines(doc)
    if fm:
        out += fm + [""]
    for _label, text, _css in doc.get("front_pages") or []:
        out += [text, ""]
    if doc["contents"]:
        out += ["CONTENTS", ""] + [p["title"] for p in doc["pieces"]] + [lbl for lbl, _t, _c in doc.get("back_pages") or []] + [""]
    for p in doc["pieces"]:
        out += ["", "* * *", "", p["title"].upper(), "", p["body"], ""]
    for label, text, _css in doc.get("back_pages") or []:
        out += ["", "* * *", "", label.upper(), "", text, ""]
    return "\n".join(out).strip("\n") + "\n"


# ---- Word -----------------------------------------------------------------

def _docx_body(document, body, style=None, like=None):
    """Add a piece's paragraphs, keeping line breaks inside a stanza. With
    `like` (a template paragraph), each new paragraph copies its style and
    its first run's font."""
    import copy

    paras = []
    for block in _blocks(body):
        p = document.add_paragraph(style=style)
        if like is not None:
            p.style = like.style
            p.paragraph_format.alignment = like.paragraph_format.alignment
        run = p.add_run()
        if like is not None and like.runs:
            rpr = like.runs[0]._r.rPr
            if rpr is not None:
                run._r.insert(0, copy.deepcopy(rpr))
        for i, line in enumerate(block):
            if i:
                run.add_break()
            run.add_text(line)
        paras.append(p)
    return paras


def _docx_page(section, trim, margin):
    from docx.shared import Inches
    _, w, h = TRIM_SIZES[trim]
    section.page_width, section.page_height = Inches(w), Inches(h)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Inches(margin))


def _margin_for(trim):
    return 1.0 if TRIM_SIZES[trim][1] >= 8 else 0.75


def _docx_cover(document, image_path, trim, new_section):
    from docx.enum.section import WD_SECTION
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches
    from PIL import Image

    section = document.add_section(WD_SECTION.NEW_PAGE) if new_section else document.sections[0]
    _docx_page(section, trim, 0)
    _, w, h = TRIM_SIZES[trim]
    with Image.open(image_path) as im:
        iw, ih = im.size
    scale = min(w / iw, (h - 0.05) / ih)
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = p.paragraph_format.space_after = 0
    p.add_run().add_picture(image_path, width=Inches(iw * scale), height=Inches(ih * scale))
    return section


def _as_docx(doc, path, template):
    import docx
    from docx.enum.section import WD_SECTION
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.shared import Pt

    if template:
        _docx_from_template(doc, path, template)
        return
    document = docx.Document()
    normal = document.styles["Normal"]
    normal.font.name = "Georgia"
    normal.font.size = Pt(11)
    trim = doc["trim"]
    margin = _margin_for(trim)

    if doc["single"]:
        _docx_page(document.sections[0], trim, margin)
        p = doc["pieces"][0]
        document.add_heading(p["title"], level=1)
        _docx_body(document, p["body"])
        document.save(path)
        return

    started = False
    if doc.get("front"):
        _docx_cover(document, doc["front"], trim, new_section=False)
        started = True

    def body_section():
        nonlocal started
        section = document.add_section(WD_SECTION.NEW_PAGE) if started else document.sections[0]
        started = True
        _docx_page(section, trim, margin)

    def page_break():
        document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    body_section()
    if doc["title_page"]:
        t = document.add_paragraph()
        t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        t.paragraph_format.space_before = Pt(144)
        r = t.add_run(doc["title"])
        r.font.size = Pt(28)
        if doc["subtitle"]:
            s = document.add_paragraph()
            s.alignment = WD_ALIGN_PARAGRAPH.CENTER
            s.add_run(doc["subtitle"]).font.size = Pt(15)
        if doc["author"]:
            a = document.add_paragraph()
            a.alignment = WD_ALIGN_PARAGRAPH.CENTER
            a.paragraph_format.space_before = Pt(48)
            a.add_run(doc["author"]).font.size = Pt(14)
        page_break()
    fm = _front_matter_lines(doc)
    if fm:
        for line in fm:
            document.add_paragraph(line).paragraph_format.space_after = Pt(2)
        page_break()
    for _label, text, _css in doc.get("front_pages") or []:
        paras = _docx_body(document, text)
        for para in paras:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in para.runs:
                run.italic = True
        if paras:
            paras[0].paragraph_format.space_before = Pt(144)
        page_break()
    back = doc.get("back_pages") or []
    if doc["contents"]:
        document.add_heading("Contents", level=1)
        for title in [p["title"] for p in doc["pieces"]] + [lbl for lbl, _t, _c in back]:
            document.add_paragraph(title).paragraph_format.space_after = Pt(2)
        page_break()
    chapters = [(p["title"], p["body"]) for p in doc["pieces"]] + [(lbl, text) for lbl, text, _c in back]
    for i, (title, body) in enumerate(chapters):
        document.add_heading(title, level=1)
        _docx_body(document, body)
        if i < len(chapters) - 1:
            page_break()
    if doc.get("back"):
        _docx_cover(document, doc["back"], trim, new_section=True)
    document.save(path)


PLACEHOLDER_RE = re.compile(r"\{\{\s*(title|body|category|date|author)\s*\}\}", re.IGNORECASE)


def _docx_from_template(doc, path, template):
    """Fill {{title}}, {{body}}, {{category}}, {{date}} in a Word template.
    Without a {{body}} placeholder, the text goes at the end."""
    import docx

    document = docx.Document(template["path"])
    piece = doc["pieces"][0]
    values = {"title": piece["title"], "category": piece.get("category", ""),
              "date": datetime.date.today().strftime("%B %d, %Y").replace(" 0", " "),
              "author": doc.get("author", "")}

    def paragraphs():
        yield from document.paragraphs
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    yield from cell.paragraphs

    body_at = None
    for p in list(paragraphs()):
        if not PLACEHOLDER_RE.search(p.text):
            continue
        if re.search(r"\{\{\s*body\s*\}\}", p.text, re.IGNORECASE) and body_at is None:
            body_at = p
            continue
        text = PLACEHOLDER_RE.sub(lambda m: values[m.group(1).lower()], p.text)
        for r in p.runs[1:]:
            r.text = ""
        if p.runs:
            p.runs[0].text = text
    if body_at is None:
        _docx_body(document, piece["body"])
    else:
        new = _docx_body(document, piece["body"], like=body_at)
        anchor = body_at._p
        for p in new:
            anchor.addnext(p._p)
            anchor = p._p
        body_at._p.getparent().remove(body_at._p)
    document.save(path)


# ---- PDF ------------------------------------------------------------------

_FONTS = None


def _fonts():
    """Register a serif family with real Unicode coverage: Georgia on
    Windows, else another installed serif, else the sans font bundled with
    ReportLab."""
    global _FONTS
    if _FONTS:
        return _FONTS
    from reportlab.lib.fonts import addMapping
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    import reportlab

    windir = os.path.join(os.environ.get("WINDIR", r"C:\Windows"), "Fonts")
    rl = os.path.join(os.path.dirname(reportlab.__file__), "fonts")
    candidates = [
        [os.path.join(windir, f) for f in ("georgia.ttf", "georgiab.ttf", "georgiai.ttf", "georgiaz.ttf")],
        [os.path.join(windir, f) for f in ("times.ttf", "timesbd.ttf", "timesi.ttf", "timesbi.ttf")],
        ["/usr/share/fonts/truetype/dejavu/" + f for f in
         ("DejaVuSerif.ttf", "DejaVuSerif-Bold.ttf", "DejaVuSerif-Italic.ttf", "DejaVuSerif-BoldItalic.ttf")],
        [os.path.join(rl, f) for f in ("Vera.ttf", "VeraBd.ttf", "VeraIt.ttf", "VeraBI.ttf")],
    ]
    for files in candidates:
        if all(os.path.isfile(f) for f in files):
            names = ["MinervaSerif", "MinervaSerif-Bold", "MinervaSerif-Italic", "MinervaSerif-BoldItalic"]
            for n, f in zip(names, files):
                pdfmetrics.registerFont(TTFont(n, f))
            addMapping("MinervaSerif", 0, 0, names[0])
            addMapping("MinervaSerif", 1, 0, names[1])
            addMapping("MinervaSerif", 0, 1, names[2])
            addMapping("MinervaSerif", 1, 1, names[3])
            _FONTS = {"regular": names[0], "bold": names[1], "italic": names[2], "file": files[0]}
            return _FONTS
    _FONTS = {"regular": "Times-Roman", "bold": "Times-Bold", "italic": "Times-Italic", "file": None}
    return _FONTS


def _pdf_markup(line):
    """One line of writing as ReportLab paragraph markup, keeping leading
    spaces (indented lines in poems)."""
    stripped = line.lstrip(" ")
    return "&nbsp;" * (len(line) - len(stripped)) + escape(stripped)


def _pdf_styles(trim):
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.styles import ParagraphStyle

    f = _fonts()
    small = TRIM_SIZES[trim][1] < 7
    size = 10.5 if small else 11.5
    return {
        "body": ParagraphStyle("body", fontName=f["regular"], fontSize=size, leading=size * 1.45,
                               spaceAfter=size * 0.9),
        "heading": ParagraphStyle("heading", fontName=f["bold"], fontSize=size * 1.6,
                                  leading=size * 2, spaceAfter=size * 1.4),
        "title": ParagraphStyle("title", fontName=f["bold"], fontSize=size * 2.6, leading=size * 3.1,
                                alignment=TA_CENTER, spaceBefore=size * 10),
        "subtitle": ParagraphStyle("subtitle", fontName=f["italic"], fontSize=size * 1.3,
                                   leading=size * 1.8, alignment=TA_CENTER, spaceBefore=size),
        "author": ParagraphStyle("author", fontName=f["regular"], fontSize=size * 1.2,
                                 leading=size * 1.6, alignment=TA_CENTER, spaceBefore=size * 4),
        "small": ParagraphStyle("small", fontName=f["regular"], fontSize=size * 0.85,
                                leading=size * 1.25),
        "toc": ParagraphStyle("toc", fontName=f["regular"], fontSize=size, leading=size * 1.6),
        "centered": ParagraphStyle("centered", fontName=f["italic"], fontSize=size, leading=size * 1.45,
                                   spaceAfter=size * 0.9, alignment=TA_CENTER),
    }


def _pdf_body(body, style):
    from reportlab.platypus import Paragraph
    return [Paragraph("<br/>".join(_pdf_markup(ln) for ln in block), style) for block in _blocks(body)]


def _as_pdf(doc, path, template):
    from reportlab.lib.utils import ImageReader
    from reportlab.platypus import (BaseDocTemplate, Frame, NextPageTemplate, PageBreak,
                                    PageTemplate, Paragraph, Spacer)

    f = _fonts()
    trim = doc["trim"]
    _, w_in, h_in = TRIM_SIZES[trim]
    page = (w_in * INCH, h_in * INCH)
    margin = _margin_for(trim) * INCH
    background = None
    if template:
        with open(template["path"], "rb") as fh:
            bw, bh, png = _pdf_background(fh.read())
        page = (bw, bh)
        background = ImageReader(io.BytesIO(png))
        margin = 1.0 * INCH
    styles = _pdf_styles(trim)
    w, h = page

    def draw_background(canvas, _doc):
        if background is not None:
            canvas.drawImage(background, 0, 0, width=w, height=h)

    def draw_number(canvas, d):
        draw_background(canvas, d)
        canvas.setFont(f["regular"], 9)
        canvas.drawCentredString(w / 2, margin / 2, str(canvas.getPageNumber()))

    def draw_cover(image_path):
        def draw(canvas, _doc):
            reader = ImageReader(image_path)
            iw, ih = reader.getSize()
            scale = min(w / iw, h / ih)
            dw, dh = iw * scale, ih * scale
            canvas.drawImage(reader, (w - dw) / 2, (h - dh) / 2, width=dw, height=dh)
        return draw

    def frame(name):
        return Frame(margin, margin, w - 2 * margin, h - 2 * margin, id=name)

    pdf = BaseDocTemplate(path, pagesize=page, title=doc["title"], author=doc.get("author") or "",
                          creator="Minerva", leftMargin=margin, rightMargin=margin,
                          topMargin=margin, bottomMargin=margin)
    if doc["single"]:
        pdf.addPageTemplates([PageTemplate("front", [frame("f")], onPage=draw_background)])
        p = doc["pieces"][0]
        pdf.build([Paragraph(escape(p["title"]), styles["heading"])] + _pdf_body(p["body"], styles["body"]))
        return

    # Front matter: one page each for the title, the copyright lines and
    # the contents, without page numbers.
    sections = []
    if doc["title_page"]:
        # spaceBefore is dropped at the top of a page, so push the title down
        sec = [Spacer(1, (h - 2 * margin) * 0.28), Paragraph(escape(doc["title"]), styles["title"])]
        if doc["subtitle"]:
            sec.append(Paragraph(escape(doc["subtitle"]), styles["subtitle"]))
        if doc["author"]:
            sec.append(Paragraph(escape(doc["author"]), styles["author"]))
        sections.append(sec)
    fm = _front_matter_lines(doc)
    if fm:
        lines = [Paragraph(escape(line) or "&nbsp;", styles["small"]) for line in fm]
        sections.append([Spacer(1, max(0, h - 2 * margin - (len(fm) + 2) * styles["small"].leading))] + lines)
    for _label, text, _css in doc.get("front_pages") or []:
        sections.append([Spacer(1, (h - 2 * margin) * 0.3)] + _pdf_body(text, styles["centered"]))
    back = doc.get("back_pages") or []
    if doc["contents"]:
        sections.append([Paragraph("Contents", styles["heading"])] +
                        [Paragraph(escape(t), styles["toc"]) for t in [p["title"] for p in doc["pieces"]]
                         + [lbl for lbl, _t, _c in back]])

    after_cover = "front" if sections else "body"
    order = []
    story = []
    if doc.get("front"):
        order.append(PageTemplate("cover", [frame("c")], onPage=draw_cover(doc["front"])))
        story += [NextPageTemplate(after_cover), Spacer(1, 1), PageBreak()]
    for i, sec in enumerate(sections):
        story += sec
        if i == len(sections) - 1:
            story.append(NextPageTemplate("body"))
        story.append(PageBreak())
    chapters = [(p["title"], p["body"]) for p in doc["pieces"]] + [(lbl, text) for lbl, text, _c in back]
    for i, (title, body) in enumerate(chapters):
        story.append(Paragraph(escape(title), styles["heading"]))
        story += _pdf_body(body, styles["body"])
        if i < len(chapters) - 1:
            story.append(PageBreak())
    if doc.get("back"):
        story += [NextPageTemplate("back"), PageBreak(), Spacer(1, 1)]

    front = PageTemplate("front", [frame("f")], onPage=draw_background)
    body = PageTemplate("body", [frame("b")], onPage=draw_number)
    order += [front, body] if sections else [body, front]
    if doc.get("back"):
        order.append(PageTemplate("back", [frame("k")], onPage=draw_cover(doc["back"])))
    pdf.addPageTemplates(order)
    pdf.build(story)
