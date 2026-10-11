import io
import zipfile

import docx
import pytest
from reportlab.pdfgen import canvas

from minerva import core, fileimport


def _docx(paras, title=None, table=False):
    d = docx.Document()
    if title:
        d.add_heading(title, level=1)
    for p in paras:
        d.add_paragraph(p)
    if table:
        d.add_table(rows=1, cols=1)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_word_poem_and_prose():
    poem = fileimport.convert("salt.docx", _docx(["Salt on the sill,", "salt in the hymn.", "", "Second stanza"], "Salt Hymn"))
    assert poem["title"] == "Salt Hymn"
    assert poem["text"] == "Salt on the sill,\nsalt in the hymn.\n\nSecond stanza\n"
    long = "This is a long paragraph of prose that goes on for a while so it reads like a story. " * 2
    prose = fileimport.convert("My Story.docx", _docx([long, long], table=True))
    assert prose["title"] == "My Story"
    assert prose["text"].count("\n\n") == 1 and "1 table(s) left out" in prose["notes"]
    with pytest.raises(fileimport.ImportFileError):
        fileimport.convert("bad.docx", b"not a zip")


def test_markdown_with_front_matter_and_links():
    md = ("---\ntitle: Harbor Note\ntags: [sea, draft]\n---\n# Ignored Heading\n"
          "The **tide** comes in  \nand _the_ [[Gulls|gulls]] [argue](http://x).\n\n![pic](a.png)\nEnd\\\n").encode()
    r = fileimport.convert("note.md", md)
    assert r["title"] == "Harbor Note" and r["fields"]["tags"] == ["sea", "draft"]
    assert r["text"] == "The tide comes in\nand the gulls argue.\n\nEnd\n"
    assert "images left out" in r["notes"]
    plain = fileimport.convert("Gull Psalm.md", b"# Gull Psalm\n\nwings\n")
    assert plain["title"] == "Gull Psalm" and plain["text"] == "wings\n"


def test_rtf_odt_html_txt():
    rtf = rb"{\rtf1\ansi{\fonttbl{\f0 Times;}}\f0 Caf\'e9 line\par second\line third \u8212? end\par}"
    assert fileimport.convert("x.rtf", rtf)["text"] == "Café line\nsecond\nthird — end\n"
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("content.xml", '<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
                   'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"><office:body><office:text>'
                   '<text:h>Night Ferry</text:h><text:p>lights<text:line-break/>on the water</text:p>'
                   '<text:p/><text:p>end</text:p></office:text></office:body></office:document-content>')
    odt = fileimport.convert("f.odt", buf.getvalue())
    assert odt["title"] == "Night Ferry" and odt["text"] == "lights\non the water\n\nend\n"
    page = b"<html><head><title>Cold Front</title><style>p{}</style></head><body><p>one<br>two</p><p>three</p></body></html>"
    h = fileimport.convert("c.html", page)
    assert h["title"] == "Cold Front" and h["text"] == "one\ntwo\nthree\n"
    assert fileimport.convert("w.txt", "naïve".encode("cp1252"))["text"] == "naïve\n"
    with pytest.raises(fileimport.ImportFileError):
        fileimport.convert("x.exe", b"MZ")


def test_pdf_text_and_scanned(tmp_path):
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.drawString(72, 720, "The keeper climbs the stairs every night.")
    c.save()
    r = fileimport.convert("Keeper.pdf", buf.getvalue())
    assert "keeper climbs" in r["text"] and r["title"] == "Keeper"
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.rect(10, 10, 100, 100)
    c.save()
    with pytest.raises(fileimport.ImportFileError, match="Scans"):
        fileimport.convert("scan.pdf", buf.getvalue())


def test_preview_route_and_create_with_tags(client, library):
    core.configure(library)
    core.scan_and_sync()
    client.post("/api/setup/library", json={"mode": "open", "path": library})
    files = {"files": [(io.BytesIO(b"# Morning Tide\n\nagain\n"), "Morning Tide.md"), (io.BytesIO(b"x"), "y.exe")]}
    out = client.post("/api/import/files/preview", data=files, content_type="multipart/form-data").get_json()
    assert out[0]["title"] == "Morning Tide" and out[0]["existing_id"] and "error" in out[1]
    made = client.post("/api/entries/new", json={"items": [{"title": "Tagged", "text": "x\n", "tags": ["sea"]}],
                                                 "category": "Poetry"}).get_json()
    assert client.get(f"/api/entries/{made[0]['id']}").get_json()["tags"] == ["sea"]


def test_a_title_line_isnt_repeated_in_the_draft():
    r = fileimport.convert("Salt Hymn.txt", b"Salt Hymn\n\nThe salt stays on the window\nlong after the wind.\n")
    assert r["title"] == "Salt Hymn" and r["text"] == "The salt stays on the window\nlong after the wind.\n"
    # a first line that's part of the poem stays
    r = fileimport.convert("Salt.txt", b"Salt\nis what the window keeps\n")
    assert r["text"].startswith("Salt\nis what")
    r = fileimport.convert("Salt.txt", b"Salt\n")
    assert r["text"] == "Salt\n"
