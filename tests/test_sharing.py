import io
import os

import docx
import pypdfium2 as pdfium
import pytest
from PIL import Image

from minerva import core, sharing

from conftest import sha, write

SEP = "=" * 60


@pytest.fixture
def lib(library):
    write(os.path.join(library, "Poetry", "Salt Hymn", "Salt Hymn [C].txt"),
          "Salt on the sill—\n  salt in the “hymn”.\n\nA second stanza\nends here.\n")
    write(os.path.join(library, "Fiction", "The Lighthouse - [LH]", "The Lighthouse [J].txt"),
          "Source: Blue Notebook Journal\nSource: Entry 4\nThe Lighthouse\nCategory: Fiction\nDate: none\n"
          + SEP + "\n\nOnce there was a keeper.\n")
    core.configure(library)
    core.scan_and_sync()
    conn = core.get_conn()
    yield library, conn
    conn.close()


def _id(conn, title):
    return conn.execute("SELECT id FROM entries WHERE title = ?", (title,)).fetchone()["id"]


def _image(fmt="JPEG"):
    buf = io.BytesIO()
    Image.new("RGB", (600, 900), (30, 50, 110)).save(buf, fmt)
    return buf.getvalue()


def test_marking_final_keeps_a_copy_and_archives_the_old_one(lib):
    library, conn = lib
    eid = _id(conn, "Salt Hymn")
    draft = os.path.join(library, "Poetry", "Salt Hymn", "Salt Hymn [C].txt")
    before = sha(draft)
    f = sharing.mark_final(conn, eid)
    assert f["name"] == "Salt Hymn (final).txt" and f["same_as_current"]
    core.save_entry_content(conn, eid, "A new version.\n", "overwrite")
    assert sharing.get_final(core.get_entry_row(conn, eid))["same_as_current"] is False

    sharing.mark_final(conn, eid)
    archive = os.listdir(os.path.join(library, "Poetry", "Salt Hymn", "Archive"))
    assert archive == ["Salt Hymn (final until %s).txt" % __import__("datetime").date.today().isoformat()]
    assert open(os.path.join(library, "Poetry", "Salt Hymn", "Archive", archive[0]), encoding="utf-8").read().startswith("Salt on")
    # the entry's own draft is still the one draft Minerva reads
    core.scan_and_sync()
    assert core.get_entry_row(conn, eid)["draft_filename"] == "Salt Hymn [C].txt"
    assert before != sha(draft)  # it was edited above, not by marking

    sharing.unmark_final(conn, eid)
    assert sharing.get_final(core.get_entry_row(conn, eid)) is None
    assert len(os.listdir(os.path.join(library, "Poetry", "Salt Hymn", "Archive"))) == 2


def test_strip_header_drops_only_a_transcription_header():
    text = "Source: X Journal\nSource: Entry 1\nTitle\nCategory: Poetry\nDate: none\n" + SEP + "\n\nBody\n"
    assert sharing.strip_header(text) == "Body"
    assert sharing.strip_header("Just a poem\n" + SEP + "\nwith a rule") == "Just a poem\n" + SEP + "\nwith a rule"


def test_entry_exports_in_every_format_without_overwriting(lib):
    library, conn = lib
    eid = _id(conn, "The Lighthouse")
    names = [sharing.export_entry(conn, eid, f)["name"] for f in sharing.FORMATS]
    assert names == ["The Lighthouse.txt", "The Lighthouse.docx", "The Lighthouse.pdf"]
    out = os.path.join(library, "Exports")
    assert open(os.path.join(out, names[0]), encoding="utf-8").read() == "The Lighthouse\n\nOnce there was a keeper.\n"
    assert [p.text for p in docx.Document(os.path.join(out, names[1])).paragraphs][-1] == "Once there was a keeper."
    assert len(pdfium.PdfDocument(os.path.join(out, names[2]))) == 1
    assert sharing.export_entry(conn, eid, "txt")["name"] == "The Lighthouse (2).txt"


def test_word_template_fills_placeholders_and_keeps_stanzas(lib):
    _, conn = lib
    d = docx.Document()
    d.add_paragraph("{{title}}")
    d.add_paragraph("{{ body }}")
    d.add_paragraph("Filed under {{category}}")
    buf = io.BytesIO()
    d.save(buf)
    t = sharing.add_template("Letter.docx", buf.getvalue(), name="Letter")
    assert [x["name"] for x in sharing.list_templates()] == ["Letter"]
    r = sharing.export_entry(conn, _id(conn, "Salt Hymn"), "docx", template_id=t["id"])
    texts = [p.text for p in docx.Document(r["path"]).paragraphs]
    assert texts == ["Salt Hymn", "Salt on the sill—\n  salt in the “hymn”.",
                     "A second stanza\nends here.", "Filed under Poetry"]
    with pytest.raises(sharing.ShareError):
        sharing.export_entry(conn, _id(conn, "Salt Hymn"), "pdf", template_id=t["id"])
    sharing.remove_template(t["id"])
    assert sharing.list_templates() == []


def test_pdf_template_sets_the_page(lib):
    _, conn = lib
    buf = io.BytesIO()
    Image.new("RGB", (595, 842), "white").save(buf, "PDF")
    t = sharing.add_template("Letterhead.pdf", buf.getvalue())
    r = sharing.export_entry(conn, _id(conn, "Salt Hymn"), "pdf", template_id=t["id"])
    w, h = pdfium.PdfDocument(r["path"])[0].get_size()
    assert (round(w), round(h)) == (595, 842)
    with pytest.raises(sharing.ShareError):
        sharing.add_template("notes.txt", b"hello")


def test_project_proof_has_covers_front_matter_and_order(lib):
    library, conn = lib
    pid = core.list_projects(conn)[0]["id"] if core.list_projects(conn) else None
    if pid is None:
        core.create_project(conn, "Lighthouse", "LH")
        pid = core.list_projects(conn)[0]["id"]
    core.set_entry_project(conn, _id(conn, "Salt Hymn"), "LH")
    core.set_project_order(conn, pid, [_id(conn, "Salt Hymn"), _id(conn, "The Lighthouse")])
    sharing.mark_final(conn, _id(conn, "Salt Hymn"))
    sharing.set_cover(pid, "front", "front.jpg", _image())
    sharing.set_cover(pid, "back", "back.png", _image("PNG"))
    s = sharing.save_project_settings(conn, pid, {"trim": "5x8", "author": "A. Writer",
                                                  "isbn": "978-0-00-000000-0", "copyright": "© 2026 A. Writer"})
    assert s["covers"] == {"front": True, "back": True} and s["trim"] == "5x8"

    r = sharing.export_project(conn, pid, "pdf")
    assert (r["entries"], r["finals"]) == (2, 1) and r["name"] == "Lighthouse - proof (5x8).pdf"
    pdf = pdfium.PdfDocument(r["path"])
    # cover, title, copyright, contents, two entries, back cover
    assert len(pdf) == 7
    assert tuple(round(x) for x in pdf[0].get_size()) == (360, 576)

    text = open(sharing.export_project(conn, pid, "txt")["path"], encoding="utf-8").read()
    assert text.index("SALT HYMN") < text.index("THE LIGHTHOUSE")
    assert "ISBN 978-0-00-000000-0" in text and "A. Writer" in text

    d = docx.Document(sharing.export_project(conn, pid, "docx")["path"])
    assert len(d.inline_shapes) == 2 and len(d.sections) == 3

    with pytest.raises(sharing.ShareError):
        sharing.save_project_settings(conn, pid, {"trim": "huge"})
    with pytest.raises(sharing.ShareError):
        sharing.set_cover(pid, "front", "cover.gif", b"GIF89a")


def test_share_routes(client, lib):
    library, conn = lib
    eid = _id(conn, "Salt Hymn")
    assert client.post(f"/api/entries/{eid}/final").status_code == 200
    assert client.get(f"/api/entries/{eid}").get_json()["final"]["name"] == "Salt Hymn (final).txt"
    rows = client.get("/api/entries?search=Salt").get_json()["rows"]
    assert rows[0]["has_final"] is True
    r = client.post(f"/api/entries/{eid}/export", json={"format": "pdf"}).get_json()
    assert r == {"name": "Salt Hymn.pdf", "from_final": True, "folder": "Exports"}
    assert client.get("/api/exports/file?name=Salt Hymn.pdf").data[:4] == b"%PDF"
    assert client.get("/api/exports/file?name=../Poetry/Salt Hymn/Salt Hymn [C].txt").status_code == 404
    up = client.post("/api/templates", data={"file": (io.BytesIO(b"not a pdf"), "x.pdf")})
    assert up.status_code == 400
    opts = client.get("/api/share/options").get_json()
    assert opts["formats"] == ["txt", "docx", "pdf"] and opts["templates"] == []


def _project(conn):
    core.create_project(conn, "Tide Book", "TB")
    pid = next(p["id"] for p in core.list_projects(conn) if p["name"] == "Tide Book")
    core.set_entry_project(conn, _id(conn, "Salt Hymn"), "TB")
    core.set_entry_project(conn, _id(conn, "The Lighthouse"), "TB")
    core.set_project_order(conn, pid, [_id(conn, "Salt Hymn"), _id(conn, "The Lighthouse")])
    return pid


def test_front_and_back_pages(lib):
    library, conn = lib
    pid = _project(conn)
    sharing.save_project_settings(conn, pid, {"trim": "5x8", "dedication": "for M.", "epigraph": "",
                                              "acknowledgments": "Thanks to Harbor Review.\r\n",
                                              "about": "A. Writer lives by the sea."})
    text = open(sharing.export_project(conn, pid, "txt")["path"], encoding="utf-8").read()
    assert text.index("for M.") < text.index("CONTENTS") < text.index("SALT HYMN") < text.index("ACKNOWLEDGMENTS") \
        < text.index("ABOUT THE AUTHOR")
    assert "Epigraph" not in text
    pdf = pdfium.PdfDocument(sharing.export_project(conn, pid, "pdf")["path"])
    # title, dedication, contents, two entries, acknowledgments, about
    assert len(pdf) == 7
    d = docx.Document(sharing.export_project(conn, pid, "docx")["path"])
    heads = [p.text for p in d.paragraphs if p.style.name.startswith("Heading")]
    assert heads == ["Contents", "Salt Hymn", "The Lighthouse", "Acknowledgments", "About the Author"]


def test_epub(lib):
    import zipfile
    import xml.etree.ElementTree as ET

    library, conn = lib
    pid = _project(conn)
    sharing.set_cover(pid, "front", "front.jpg", _image())
    sharing.save_project_settings(conn, pid, {"author": "A. Writer", "isbn": "978-0-00-000000-0",
                                              "copyright": "© 2026 A. Writer", "acknowledgments": "Thanks."})
    r = sharing.export_project(conn, pid, "epub")
    assert r["name"] == "Tide Book.epub"
    z = zipfile.ZipFile(r["path"])
    first = z.infolist()[0]
    assert first.filename == "mimetype" and first.compress_type == zipfile.ZIP_STORED
    assert z.read("mimetype") == b"application/epub+zip"
    for name in z.namelist():
        if name.endswith((".xhtml", ".opf", ".ncx", ".xml")):
            ET.fromstring(z.read(name))  # every document is well-formed XML
    opf = z.read("OEBPS/content.opf").decode()
    assert "urn:isbn:9780000000000" in opf and "<dc:creator>A. Writer</dc:creator>" in opf
    assert 'properties="cover-image"' in opf and 'properties="nav"' in opf
    ns = {"o": "http://www.idpf.org/2007/opf"}
    root = ET.fromstring(opf)
    hrefs = {i.get("id"): i.get("href") for i in root.find("o:manifest", ns)}
    spine = [hrefs[i.get("idref")] for i in root.find("o:spine", ns)]
    titles = [ET.fromstring(z.read("OEBPS/" + h)).find(".//{http://www.w3.org/1999/xhtml}title").text for h in spine]
    assert titles == ["Cover", "Tide Book", "Copyright", "Contents", "Salt Hymn", "The Lighthouse", "Acknowledgments"]
    hymn = z.read("OEBPS/" + spine[4]).decode()
    # line breaks and indents kept
    assert "Salt on the sill—<br/>&#160;&#160;salt in the “hymn”.</p>" in hymn


def test_project_author_starts_as_the_manuscript_name(library):
    from minerva import manuscript
    core.configure(library)
    core.scan_and_sync()
    conn = core.get_conn()
    core.create_project(conn, "Tides", "TD")
    pid = core.list_projects(conn)[0]["id"]
    assert sharing.project_settings(conn, pid)["author"] == ""
    manuscript.save_contact({"name": "A. Writer"})
    assert sharing.project_settings(conn, pid)["author"] == "A. Writer"
    sharing.save_project_settings(conn, pid, {"author": ""})  # cleared on purpose stays cleared
    assert sharing.project_settings(conn, pid)["author"] == ""
    conn.close()
