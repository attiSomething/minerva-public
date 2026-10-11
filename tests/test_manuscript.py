import os

import docx
import pypdfium2 as pdfium
import pytest

from minerva import core, manuscript, sharing, submissions

from conftest import write

SEP = "=" * 60


@pytest.fixture
def lib(library):
    write(os.path.join(library, "Poetry", "Salt Hymn", "Salt Hymn [C].txt"),
          "Salt on the sill,\n  salt in the hymn.\n\nA second stanza\nends here.\n\n\nThird.\n")
    write(os.path.join(library, "Fiction", "The Lighthouse - [LH]", "The Lighthouse [J].txt"),
          "Source: Blue Notebook Journal\nSource: Entry 4\nThe Lighthouse\nCategory: Fiction\nDate: none\n"
          + SEP + "\n\nOnce there was a keeper.\n")
    core.configure(library)
    core.scan_and_sync()
    conn = core.get_conn()
    yield library, conn
    conn.close()


def _row(conn, title):
    return conn.execute("SELECT * FROM entries WHERE title = ?", (title,)).fetchone()


def test_line_and_stanza_counts_skip_the_header(lib):
    _, conn = lib
    salt = _row(conn, "Salt Hymn")
    assert (salt["line_count"], salt["stanza_count"]) == (5, 3)
    keeper = _row(conn, "The Lighthouse")
    assert (keeper["line_count"], keeper["stanza_count"]) == (1, 1)
    # saving updates them
    core.save_entry_content(conn, salt["id"], "One line only.\n", "overwrite")
    salt = _row(conn, "Salt Hymn")
    assert (salt["line_count"], salt["stanza_count"]) == (1, 1)


def test_contact_details_are_kept_outside_the_library(lib):
    library, _ = lib
    saved = manuscript.save_contact({"name": "  Ada   Quill ", "address": "1 Harbor Rd\r\nPort Town, ME", "email": "ada@example.com"})
    assert saved["name"] == "Ada Quill" and saved["address"] == "1 Harbor Rd\nPort Town, ME"
    assert manuscript.contact()["email"] == "ada@example.com"
    for d, _, names in os.walk(library):
        for n in names:
            assert "ada@example.com" not in open(os.path.join(d, n), encoding="utf-8", errors="ignore").read()
    assert manuscript.contact_lines(manuscript.contact()) == ["Ada Quill", "1 Harbor Rd", "Port Town, ME", "ada@example.com"]
    assert manuscript.running_head("Ada Quill", "Salt Hymn") == "Quill / SALT HYMN / "


def test_entry_exports_in_standard_manuscript(lib):
    _, conn = lib
    manuscript.save_contact({"name": "Ada Quill", "address": "1 Harbor Rd", "email": "ada@example.com"})
    eid = _row(conn, "Salt Hymn")["id"]
    r = sharing.export_entry(conn, eid, "docx", "current", sharing.MANUSCRIPT)
    d = docx.Document(r["path"])
    texts = [p.text for p in d.paragraphs]
    assert texts[0] == "Ada Quill\t5 lines" and texts[1:3] == ["1 Harbor Rd", "ada@example.com"]
    assert "SALT HYMN" in texts
    stanza = [p for p in d.paragraphs if p.text.startswith("Salt on the sill")][0]
    assert stanza.paragraph_format.keep_together and "\n" in stanza.text
    assert d.sections[0].page_width.inches == pytest.approx(8.5)
    head = d.sections[0].header.paragraphs[0]
    assert head.text.startswith("Quill / SALT HYMN /") and "PAGE" in head._p.xml
    assert d.sections[0].different_first_page_header_footer

    r = sharing.export_entry(conn, eid, "pdf", "current", sharing.MANUSCRIPT)
    pdf = pdfium.PdfDocument(r["path"])
    page = pdf[0]
    assert page.get_size() == pytest.approx((612, 792))
    text = page.get_textpage().get_text_range()
    assert "Ada Quill" in text and "5 lines" in text and "SALT HYMN" in text
    pdf.close()

    r = sharing.export_entry(conn, eid, "txt", "current", sharing.MANUSCRIPT)
    txt = open(r["path"], encoding="utf-8").read()
    assert txt.startswith("Ada Quill") and "5 lines" in txt.splitlines()[0] and "SALT HYMN" in txt


def test_long_poem_gets_running_header_and_project_manuscript(lib):
    library, conn = lib
    manuscript.save_contact({"name": "Ada Quill"})
    long = "\n\n".join("\n".join(f"line {s}-{i}" for i in range(6)) for s in range(20))
    write(os.path.join(library, "Poetry", "Long Tide - [HP]", "Long Tide [C].txt"), long)
    core.create_project(conn, "Harbor Poems", "HP")
    core.scan_and_sync()
    conn = core.get_conn()
    pid = conn.execute("SELECT id FROM projects WHERE tag = 'HP'").fetchone()["id"]
    sharing.save_project_settings(conn, pid, {"trim": "manuscript", "title_page": True})
    r = sharing.export_project(conn, pid, "pdf")
    assert r["name"].startswith("Harbor Poems - manuscript")
    pdf = pdfium.PdfDocument(r["path"])
    assert len(pdf) >= 3
    first = pdf[0].get_textpage().get_text_range()
    assert "HARBOR POEMS" in first and "1 poem" in first and "1 poems" not in first
    later = pdf[2].get_textpage().get_text_range()
    assert "Quill / LONG TIDE / 3" in later
    pdf.close()
    # every stanza stays whole on a page
    for i in range(1, len(pdfium.PdfDocument(r["path"]))):
        t = pdfium.PdfDocument(r["path"])[i].get_textpage().get_text_range()
        for s in range(20):
            if f"line {s}-0" in t:
                assert f"line {s}-5" in t
    with pytest.raises(sharing.ShareError):
        sharing.save_project_settings(conn, pid, {"trim": "huge"})


def test_venue_line_limit(lib):
    _, conn = lib
    submissions.save_venue(conn, {"name": "Harbor Review", "max_lines": "40"})
    assert submissions.venue(conn, "harbor review")["max_lines"] == 40
    with pytest.raises(submissions.SubmissionError):
        submissions.save_venue(conn, {"name": "Harbor Review", "max_lines": "forty"})
    submissions.save_venue(conn, {"name": "Harbor Review", "max_lines": ""})
    assert submissions.venue(conn, "Harbor Review")["max_lines"] is None


def test_share_options_offer_the_manuscript_layout(client, lib):
    library, _ = lib
    client.post("/api/setup/library", json={"mode": "open", "path": library})
    opts = client.get("/api/share/options").get_json()
    assert opts["trims"][-1]["id"] == "manuscript" and opts["contact_set"] is False
    client.post("/api/settings/contact", json={"name": "Ada Quill"})
    assert client.get("/api/share/options").get_json()["contact_set"] is True
